"""
write_xlsx.py — formatted 4-sheet workbook deliverable
(Classification table / Statistics with live % formulas / Method notes /
AI-journal_change). Called from run_local.py when WRITE_XLSX = True.
Version 3.0 — 2026-07-28.
"""

import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AR = Font(name="Arial", size=10)
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
BORD = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")

TABLE_COLS = ["Article ID", "Reference", "MWP technique", "Relevance for TC MWP",
              "Classification category", "Application axis",
              "Test methodology axis", "Data analysis type"]


def _header(ws, headers, widths):
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, j, h)
        cell.font, cell.fill, cell.border = HDR, FILL, BORD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w


def write(arts, st, path):
    table = arts[TABLE_COLS]
    n, ng = st["n"], st["n_geo"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Classification table"
    _header(ws, TABLE_COLS, [10, 34, 14, 22, 34, 40, 44, 38])
    for i, row in enumerate(table.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.font, cell.border, cell.alignment = AR, BORD, WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(table) + 1}"

    stx = wb.create_sheet("Statistics")
    stx.column_dimensions["A"].width = 44
    stx.column_dimensions["B"].width = 12
    stx.column_dimensions["C"].width = 12
    r = 1

    def block(title, series, denom, note=""):
        nonlocal r
        c = stx.cell(r, 1, title + (f"  —  {note}" if note else ""))
        c.font = Font(name="Arial", size=10, bold=True)
        r += 1
        for h, j in (("Category", 1), ("Count", 2), ("%", 3)):
            cell = stx.cell(r, j, h)
            cell.font, cell.fill, cell.border = HDR, FILL, BORD
        r += 1
        first = r
        for k, v in series.items():
            stx.cell(r, 1, k).font = AR
            stx.cell(r, 2, int(v)).font = AR
            pc = stx.cell(r, 3, f"=B{r}/{denom}")
            pc.font, pc.number_format = AR, "0.0%"
            for j in (1, 2, 3):
                stx.cell(r, j).border = BORD
            r += 1
        stx.cell(r, 1, "Total occurrences").font = Font(name="Arial", size=10, italic=True)
        stx.cell(r, 2, f"=SUM(B{first}:B{r-1})").font = AR
        r += 2

    block(f"1. MWP technique (unit: article, n = {n})", st["technique"], n)
    block(f"2. Material types (multi-label, % of n = {n} articles)", st["materials"],
          n, "an article may count in several categories")
    block(f"3. Relevance for TC MWP (unit: article, n = {n})", st["relevance"], n)
    block(f"4. Specimen geometries in UT/IRT blocks (multi-label, % of n = {ng})",
          st["geometry"], ng, "articles with UT/IRT specimen data only")

    mn = wb.create_sheet("Method notes")
    mn.column_dimensions["A"].width = 120
    notes = [
        "OPERATIONAL DEFINITIONS — all labels derive exclusively from review-form fields (no external sources).",
        "",
        "Classification rules are stored in the six editable AI_tables/*.csv files (see README): technique family "
        "keywords, conventional-test keywords, application-axis keywords, material normalization, geometry "
        "normalization, comparison-assessed keywords. Correct any misclassification there and re-run.",
        "",
        "MWP technique — from 'C2. What type of test is being reported?' (Tests 1-4); 'Other' free text matched "
        "against AI_table_technique_keywords.csv (UT-family / IRT-family).",
        "Relevance — reviewer field 'Relevance for the laboratory characterization of bituminous materials from "
        "MWP tests' (High / Medium / Low).",
        "Classification category — conventional test reported AND comparison assessed "
        "(AI_table_comparison_keywords.csv) -> Comparative study; conventional without assessed comparison -> "
        "co-reported; else Standalone MWP study.",
        "Application axis — keyword scoring (AI_table_application_keywords.csv) over properties determined, "
        "rheological-modeling purpose and key contribution; dominant family (+ strong secondary).",
        "Test methodology axis — UT: configuration + P/S/R waves; IRT: boundary conditions + vibration modes.",
        "Data analysis type — signal domain (time / frequency / both) x resolution method (analytical / numerical "
        "inverse / empirical-regression / qualitative).",
        "Geometries — only C3.2 fields of UT/IRT test blocks; normalized via AI_table_geometry_keywords.csv; "
        "multi-label (percentages sum to more than 100%).",
        "Duplicates — P95 and P131 double-reviewed; merged by union of labels; discordant relevance flagged.",
    ]
    for i, t in enumerate(notes, 1):
        cell = mn.cell(i, 1, t)
        cell.font = Font(name="Arial", size=10, bold=t.startswith("OPERATIONAL"))
        cell.alignment = WRAP

    aj = wb.create_sheet("AI-journal_change")
    rows = []
    try:
        for line in open("AI-journal_change.md"):
            line = line.strip()
            if line.startswith("|") and not set(line) <= set("|- :"):
                cells = [x.strip() for x in line.strip("|").split("|")]
                if rows and cells[0] == "ID":   # repeated header of a later session table
                    continue
                rows.append(cells)
    except FileNotFoundError:
        rows = [["AI-journal_change.md not found"]]
    widths = [6, 12, 20, 10, 55, 45, 40]
    for j, w in enumerate(widths, 1):
        aj.column_dimensions[get_column_letter(j)].width = w
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            cell = aj.cell(i, j, re.sub(r"`|\*\*", "", v))
            cell.border, cell.alignment = BORD, WRAP
            cell.font = HDR if i == 1 else AR
            if i == 1:
                cell.fill = FILL
    aj.freeze_panes = "A2"

    wb.save(path)
