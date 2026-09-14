"""write_xlsx.py — builds the classification workbook.

Library file: no print() statements. Colours/fonts come from the XLSX config
dict in run_classify.py.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REVIEWER_NOTE = ('Reviewer Assessment is used ONLY as a tie-break criterion and '
                 'does not affect the technical contribution scores.')

MAIN_COLUMNS = [  # (header, colour key, result-frame column)
    ('Paper ID', 'grey', 'Paper ID'),
    ('Authors', 'grey', 'Authors'),
    ('Year', 'grey', 'Year'),
    ('Title', 'grey', 'Title'),
    ('MWP Technique', 'grey', 'MWP Technique'),
    ('Traditional test(s) also performed', 'grey', 'Traditional tests'),
    ('TC MWP Relevance', 'grey', 'TC MWP Relevance'),
    ('Application Contribution Score (0-3)', 'blue', 'App Score'),
    ('Application Evidence', 'blue', 'App Evidence'),
    ('Experimental Methodology Contribution Score (0-3)', 'green', 'Meth Score'),
    ('Methodology Evidence', 'green', 'Meth Evidence'),
    ('Data Analysis Contribution Score (0-3)', 'red', 'DA Score'),
    ('Signal Processing Score (0-2)', 'lred', 'SigProc Score'),
    ('Property Complexity Score (0-3)', 'lred', 'PropC Score'),
    ('Identification Method Score (0-3)', 'lred', 'IdM Score'),
    ('Signal Processing Type(s)', 'lred', 'SigProc Types'),
    ('Property / Output Type(s)', 'lred', 'Property Types'),
    ('Property Identification Method(s)', 'lred', 'Id Methods'),
    ('Raw Property Description', 'lred', 'Raw Property'),
    ('Raw Identification Method Description', 'lred', 'Raw IdM'),
    ('Rheological Modelling Score (0-2)', 'purple', 'Rheo Score'),
    ('Rheological Model(s)', 'purple', 'Rheo Models'),
    ('Rheological Modelling Purpose', 'purple', 'Rheo Purpose'),
    ('Reviewer Support Score (0-2)', 'orange', 'Reviewer Support'),
    ('ABK — Key Contribution', 'orange', 'ABK'),
    ('ABL — Reviewer Notes', 'orange', 'ABL'),
    ('Primary Contribution', 'grey', 'Primary'),
    ('Secondary Contribution', 'grey', 'Secondary'),
    ('Tie-break Used?', 'grey', 'Tie-break Used?'),
    ('Tie-break Reason', 'grey', 'Tie-break Reason'),
    ('Classification Confidence', 'grey', 'Confidence'),
    ('Material Type(s)', 'grey', 'Materials'),
    ('Specimen Geometry Type(s)', 'grey', 'Geometries'),
    ('Classification Rationale', 'grey', 'Rationale'),
]

MAIN_WIDTHS = [9, 22, 7, 40, 12, 24, 12, 10, 40, 10, 38, 10, 9, 9, 9, 26, 28, 30,
               34, 34, 9, 20, 30, 9, 36, 36, 20, 20, 9, 34, 12, 30, 26, 40]

README_LINES = [
    ('TC 331-MWP — Classification of 155 literature-review forms', True),
    ('Source: worksheet "Final for analysis AI and code" of '
     'TC_RILEM_Form_TX_AKLB_FULL VERSION.xlsx. Only form content was used; '
     'original papers and external sources were not consulted.', False),
    ('', False),
    ('IMPORTANT: ' + REVIEWER_NOTE, True),
    ('', False),
    ('All classification rules are stored in the editable AI_tables/*.csv files '
     'of the mwp_contribution module — correct the tables, not the code, then '
     'rerun run_classify.py to regenerate this workbook and all figures.', False),
    ('', False),
    ('Colour code (Main Table):', True),
    ('  Blue = Application Contribution', False),
    ('  Green = Experimental Methodology Contribution', False),
    ('  Red = Data Analysis Contribution', False),
    ('  Light red = Signal Processing / Property Complexity / Identification '
     'Method (sub-components of Data Analysis)', False),
    ('  Purple = Rheological Modelling (descriptive/supporting only)', False),
    ('  Orange = Reviewer Assessment (tie-break only)', False),
    ('', False),
    ('Scoring rules (deterministic, keyword-based on form text; '
     '"No / Not reported / N/A" never earns credit):', True),
    ('  Application 0-3: count of application evidence items (see '
     'AI_table_application_keywords.csv, plus conventional-test comparison, '
     'multiple material types and multi-temperature detection). '
     '0 items=0; 1=1; 2-3=2; >=4 or conventional validation with >=4 items=3.', False),
    ('  Methodology 0-3: count of informative test-parameter fields '
     '(IR: EF-EL; UT: GM-GS and NA-NG). 0=0; 1-2=1; 3-5=2; >=6=3.', False),
    ('  Signal Processing 0-2 (IR: EN-EP+comments; UT: GU-GW, NI-NK+comments): '
     'advanced technique or >=3 basic techniques = 2; any basic = 1; none = 0. '
     'See AI_table_signal_processing.csv.', False),
    ('  Property Complexity 0-3 (ER, GY, NM): constitutive parameters or >=2 '
     'viscoelastic properties = 3; one viscoelastic = 2; mechanical = 1; basic '
     'wave output = 0. See AI_table_property_complexity.csv.', False),
    ('  Identification Method 0-3 (IR: ES-EX; UT: GZ-HG, NN-NU): inverse/'
     'back-analysis = 3; FRF/modal/peak-based = 2; closed-form/analytical/'
     'empirical = 1; not described = 0. See AI_table_identification_method.csv.', False),
    ('  Data Analysis 0-3: 3 requires meaningful property (PropC>=1) AND '
     'IdM=3. 2 requires meaningful property AND sufficiently described '
     'processing/identification. 1 = basic/direct chain. 0 = no chain.', False),
    ('  Rheological Modelling 0-2 (AAY-ABC): model reported = 1; + TTSP / shift '
     'factors / master-curve purpose = 2. Descriptive only.', False),
    ('  Primary Contribution = highest of Application / Methodology / Data '
     'Analysis. Ties resolved by ABK/ABL emphasis only '
     '(AI_table_tiebreak_keywords.csv); if still tied, richer form evidence is '
     'used and confidence is flagged Low.', False),
    ('  Secondary Contribution reported only when the second-highest technical '
     'score is >= 2.', False),
    ('', False),
    ('When information is unclear the tables state "Unclear from the review '
     'form" / "Not reported" instead of assumptions.', False),
]


def _fills(xcfg):
    return {k: PatternFill('solid', fgColor=v) for k, v in xcfg['fills'].items()}


def write_workbook(res, empty_cols, au_lookup, out_path, xcfg):
    """Build and save the full classification workbook."""
    N = len(res)
    F = _fills(xcfg)
    font = xcfg['font']
    HFONT = Font(name=font, bold=True, size=xcfg['font_size'])
    BFONT = Font(name=font, size=xcfg['font_size'])
    WRAP = Alignment(wrap_text=True, vertical='top')
    THIN = Border(*[Side(style='thin', color=xcfg['border_color'])] * 4)

    wb = Workbook()
    wb.remove(wb.active)

    def style_header(ws, fills):
        for j, f in enumerate(fills, start=1):
            c = ws.cell(row=1, column=j)
            c.font = HFONT
            c.fill = f
            c.alignment = Alignment(wrap_text=True, vertical='center')
            c.border = THIN
        ws.freeze_panes = 'A2'

    def body_style(ws, maxr, maxc):
        for r in range(2, maxr + 1):
            for j in range(1, maxc + 1):
                c = ws.cell(row=r, column=j)
                c.font = BFONT
                c.alignment = WRAP
                c.border = THIN

    # ---- ReadMe ----
    ws = wb.create_sheet('ReadMe')
    for k, (t, b) in enumerate(README_LINES, start=1):
        c = ws.cell(row=k, column=1, value=t)
        c.font = Font(name=font, size=xcfg['font_size'], bold=b)
        c.alignment = WRAP
    ws.column_dimensions['A'].width = xcfg['readme_width']

    # ---- Main Table ----
    ws = wb.create_sheet('Main Table')
    ws.append([h for h, _, _ in MAIN_COLUMNS])
    for _, r in res.iterrows():
        row = []
        for _, _, key in MAIN_COLUMNS:
            v = r[key]
            if isinstance(v, list):
                v = '; '.join(v)
            row.append(v)
        ws.append(row)
    style_header(ws, [F[k] for _, k, _ in MAIN_COLUMNS])
    body_style(ws, N + 1, len(MAIN_COLUMNS))
    for j, w in enumerate(MAIN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.cell(row=N + 3, column=1, value=REVIEWER_NOTE).font = Font(
        name=font, bold=True, size=xcfg['font_size'], color=xcfg['warning_color'])

    # ---- Summary Tables ----
    ws = wb.create_sheet('Summary Tables')
    state = {'r': 1}

    def write_summary(title, rows, exclusive=True, note=None):
        r0 = state['r']
        ws.cell(row=r0, column=1, value=title).font = Font(
            name=font, bold=True, size=xcfg['font_size'] + 1)
        r0 += 1
        for j, h in enumerate(['Category', 'Paper Count', 'Percentage',
                               'Paper IDs'], start=1):
            c = ws.cell(row=r0, column=j, value=h)
            c.font = HFONT
            c.fill = F['grey']
            c.border = THIN
        r0 += 1
        start = r0
        for cat, pids in rows:
            ws.cell(row=r0, column=1, value=cat).font = BFONT
            ws.cell(row=r0, column=2, value=len(pids)).font = BFONT
            pc = ws.cell(row=r0, column=3, value=f'=B{r0}/{N}')
            pc.number_format = '0.0%'
            pc.font = BFONT
            idc = ws.cell(row=r0, column=4, value=', '.join(pids))
            idc.font = BFONT
            idc.alignment = WRAP
            for j in range(1, 5):
                ws.cell(row=r0, column=j).border = THIN
            r0 += 1
        if exclusive:
            ws.cell(row=r0, column=1, value='Total').font = HFONT
            t = ws.cell(row=r0, column=2, value=f'=SUM(B{start}:B{r0 - 1})')
            t.font = HFONT
            r0 += 1
        if note:
            ws.cell(row=r0, column=1, value=note).font = Font(
                name=font, italic=True, size=xcfg['font_size'] - 1)
            r0 += 1
        state['r'] = r0 + 1

    def vc(col):
        return [(k, res.loc[res[col] == k, 'Paper ID'].tolist())
                for k in res[col].value_counts().index]

    def pooled_rows(series_of_lists):
        counts = {}
        for pid, lst in zip(res['Paper ID'], series_of_lists):
            for v in lst:
                counts.setdefault(v, []).append(pid)
        return sorted(counts.items(), key=lambda kv: -len(kv[1]))

    def split_types(colname, drop):
        out = []
        for v in res[colname]:
            out.append([x.strip() for x in str(v).split(',')
                        if x.strip() and x.strip() not in drop])
        return out

    write_summary('MWP Technique (from column AU only)', vc('MWP Technique'))
    write_summary('TC MWP Relevance (column ABJ — descriptive only)',
                  vc('TC MWP Relevance'))
    write_summary('Primary Contribution (exactly one per paper, n = 155)',
                  vc('Primary'))
    write_summary('Secondary Contribution (reported only when 2nd-highest score '
                  '>= 2; not part of the main distribution)', vc('Secondary'))
    write_summary('Material Type (columns H:AT only; non-exclusive)',
                  pooled_rows(res['Materials']), exclusive=False,
                  note='A paper may report several material types; percentages '
                       'need not sum to 100%.')
    write_summary('Specimen Geometry (all "Geometry" columns; non-exclusive)',
                  pooled_rows(res['Geometries']), exclusive=False,
                  note='A paper may report several geometries; percentages need '
                       'not sum to 100%.')
    write_summary('Signal Processing Type (normalized; non-exclusive)',
                  pooled_rows(split_types(
                      'SigProc Types',
                      ('Not described', 'Described (unclassified)'))),
                  exclusive=False)
    write_summary('Property / Output Type from UT/IRT (normalized; '
                  'non-exclusive)',
                  pooled_rows(split_types(
                      'Property Types',
                      ('Not reported', 'Other output (see raw)'))),
                  exclusive=False)
    write_summary('Property Identification Method (normalized; non-exclusive)',
                  pooled_rows(split_types(
                      'Id Methods', ('Not described', 'Domain reported only'))),
                  exclusive=False)
    for j, w in zip(range(1, 5), xcfg['summary_widths']):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ---- Score Distributions ----
    ws = wb.create_sheet('Score Distributions')
    r0 = 1
    for title, colname, hi in [
            ('Application Contribution Score', 'App Score', 3),
            ('Methodology Contribution Score', 'Meth Score', 3),
            ('Data Analysis Contribution Score', 'DA Score', 3),
            ('Signal Processing Score (sub)', 'SigProc Score', 2),
            ('Property Complexity Score (sub)', 'PropC Score', 3),
            ('Identification Method Score (sub)', 'IdM Score', 3),
            ('Rheological Modelling Score', 'Rheo Score', 2),
            ('Reviewer Support Score (tie-break only)', 'Reviewer Support', 2)]:
        ws.cell(row=r0, column=1, value=title).font = Font(
            name=font, bold=True, size=xcfg['font_size'] + 1)
        r0 += 1
        for j, h in enumerate(['Score', 'Paper Count', 'Percentage'], 1):
            c = ws.cell(row=r0, column=j, value=h)
            c.font = HFONT
            c.fill = F['grey']
            c.border = THIN
        r0 += 1
        for s in range(hi + 1):
            ws.cell(row=r0, column=1, value=s).font = BFONT
            ws.cell(row=r0, column=2,
                    value=int((res[colname] == s).sum())).font = BFONT
            pc = ws.cell(row=r0, column=3, value=f'=B{r0}/{N}')
            pc.number_format = '0.0%'
            pc.font = BFONT
            for j in range(1, 4):
                ws.cell(row=r0, column=j).border = THIN
            r0 += 1
        r0 += 1
    for j, w in zip(range(1, 4), xcfg['dist_widths']):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ---- Papers without explicit UT/IRT ----
    ws = wb.create_sheet('No explicit UT-IRT')
    ws.append(['Paper ID', 'Authors', 'Year', 'Title', 'Tests reported in AU',
               'TC MWP Relevance', 'ABK — Key Contribution',
               'ABL — Reviewer Notes'])
    sub = res[res['MWP Technique'] == 'Other / not explicitly UT or IR']
    for _, r in sub.iterrows():
        ws.append([r['Paper ID'], r['Authors'], r['Year'], r['Title'],
                   au_lookup.get(r['Paper ID'], ''), r['TC MWP Relevance'],
                   r['ABK'], r['ABL']])
    style_header(ws, [F['grey']] * 8)
    body_style(ws, len(sub) + 1, 8)
    for j, w in zip(range(1, 9), xcfg['noutirt_widths']):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.cell(row=len(sub) + 3, column=1,
            value='These papers are NOT reclassified as UT or IR and are '
                  'excluded from the UT/IRT/Both denominator (see Figure 1).'
            ).font = Font(name=font, italic=True, size=xcfg['font_size'] - 1)

    # ---- Empty Columns ----
    ws = wb.create_sheet('Empty Columns')
    ws.append(['Excel column letter', 'Exact form-column title'])
    for e in empty_cols:
        ws.append(list(e))
    style_header(ws, [F['grey']] * 2)
    body_style(ws, len(empty_cols) + 1, 2)
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 110
    ws.cell(row=len(empty_cols) + 3, column=1,
            value=f'{len(empty_cols)} columns are blank for ALL {N} papers. '
                  'Cells containing "No", "Not reported" or "Not applicable" '
                  'are NOT treated as empty.').font = Font(
        name=font, italic=True, size=xcfg['font_size'] - 1)

    wb.save(out_path)
    return out_path


if __name__ == '__main__':
    raise SystemExit(
        'write_xlsx.py is a library file and cannot be run directly.\n'
        'Run the entry point instead:\n'
        '    python run_classify.py\n'
        'In Spyder: open run_classify.py, then press F5 (Run file).'
    )
